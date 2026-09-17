#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把「性格色彩测试」Excel 题库解析为结构化 JSON。

用法：
    python3 scripts/parse_questions.py <Excel 路径> [-o 输出路径]

Excel 版式说明（来自 2023-12 版《性格色彩测试.xlsx》）：
  - Sheet1，111 行 × 5 列
  - 第 1-4 行是标题 / 说明 / 答题格式提示 / 空行，跳过
  - 题目分左右两栏排版：第 1 列放第 1-15 题，第 3 列放第 16-30 题
  - 每道题形如：题干行「1. 关于人生观……」+ 四个选项行「A、……」「B、……」
  - 部分选项文本在 Excel 里折行，靠 5 个空格缩进续行，解析时合并回同一选项
  - 第 30 题特殊：每个选项是一组格言（多条），同样按续行合并

选项与色彩对应：A=红、B=蓝、C=黄、D=绿
"""
import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HEADER_ROWS = 4
QUESTION_RE = re.compile(r"^(\d+)\s*[.．、]\s*(.+)")
OPTION_RE = re.compile(r"^([A-Da-d])\s*[、,.\s]\s*(.+)")


def parse_column(rows, col_idx):
    """解析一栏题目。col_idx=0 左栏（1-15题），col_idx=2 右栏（16-30题）。"""
    result = []
    current = None
    current_opt = None

    for row in rows:
        line = row[col_idx] if col_idx < len(row) else ""
        if not line:
            continue

        q_match = QUESTION_RE.match(line)
        opt_match = OPTION_RE.match(line)

        if q_match:
            if current is not None:
                result.append(current)
            current = {
                "id": int(q_match.group(1)),
                "question": q_match.group(2).strip(),
                "options": {"A": "", "B": "", "C": "", "D": ""},
            }
            current_opt = None
        elif opt_match and current is not None:
            current_opt = opt_match.group(1).upper()
            current["options"][current_opt] = opt_match.group(2).strip()
        elif current is not None and current_opt is not None:
            # 续行：折行的选项文本，合并回当前选项
            current["options"][current_opt] += line.strip()

    if current is not None:
        result.append(current)
    return result


def main():
    ap = argparse.ArgumentParser(description="解析性格色彩测试 Excel 为 JSON")
    ap.add_argument("excel", help="Excel 文件路径")
    ap.add_argument("-o", "--output", default=None, help="输出 JSON 路径（默认 data/questions.json）")
    args = ap.parse_args()

    excel_path = Path(args.excel).expanduser()
    if not excel_path.exists():
        sys.exit(f"找不到文件：{excel_path}")

    out_path = Path(args.output).expanduser() if args.output else Path(__file__).resolve().parent.parent / "data" / "questions.json"

    ws = load_workbook(excel_path).active
    rows = []
    for r in ws.iter_rows(min_row=HEADER_ROWS + 1, values_only=True):
        rows.append([str(c).strip() if c is not None else "" for c in r])

    questions = parse_column(rows, 0) + parse_column(rows, 2)
    questions.sort(key=lambda q: q["id"])

    print(f"共解析 {len(questions)} 道题目")
    bad = 0
    for q in questions:
        missing = [k for k, v in q["options"].items() if not v]
        if missing:
            bad += 1
            print(f"  [缺失] 第 {q['id']} 题缺少选项：{missing}")
        lens = {k: len(v) for k, v in q["options"].items()}
        print(f"  第 {q['id']:>2} 题 A={lens['A']:>3}字 B={lens['B']:>3}字 C={lens['C']:>3}字 D={lens['D']:>3}字  {q['question'][:24]}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(questions, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"\n已写入 {out_path}（{len(questions)} 题，异常 {bad} 题）")
    return 0 if bad == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
